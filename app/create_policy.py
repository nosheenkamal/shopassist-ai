from docx import Document

doc = Document()
doc.add_heading("ShopAssist AI — Demo Refund Policy", 0)
doc.add_paragraph("CAPSTONE DEMONSTRATION DOCUMENT")

doc.add_heading("Purpose", level=1)
doc.add_paragraph(
    "This sample policy is for the ShopAssist AI educational capstone. "
    "It is not a verified policy for a real store."
)

doc.add_heading("Demo Eligibility Rules", level=1)
rules = [
    "A customer may request a return or refund within 7 calendar days of delivery.",
    "The item must be unused and in its original packaging.",
    "The customer must provide an order ID so the order can be checked.",
    "Meeting these example rules does not guarantee approval. A support representative must review the request.",
    "If the order or delivery date cannot be verified, mark the request as Pending Review."
]
for rule in rules:
    doc.add_paragraph(rule, style="List Bullet")

doc.add_heading("Demo Refund Workflow", level=1)
steps = [
    "Find the order in the Orders spreadsheet.",
    "Check the order details and delivery date, if available.",
    "Compare the request with the demo eligibility rules.",
    "Record the result as Approved for Demo, Rejected, or Pending Review.",
    "Send a confirmation email only after an authorized reviewer approves the refund."
]
for step in steps:
    doc.add_paragraph(step, style="List Number")

doc.add_heading("Important Notice", level=1)
doc.add_paragraph(
    "This is sample data for an educational capstone. "
    "It does not authorize real refunds, payments, or customer communications."
)

doc.save("refund_policy.docx")
print("Success! Created refund_policy.docx")